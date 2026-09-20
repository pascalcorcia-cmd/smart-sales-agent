import io

from pptx import Presentation
from pptx.util import Inches, Pt

from app.markdown_blocks import parse_blocks, split_bold

MAX_BULLETS_PER_SLIDE = 8


def _set_bullet_text(paragraph, text: str, level: int):
    paragraph.level = level
    paragraph.font.size = Pt(16 if level else 20)
    for segment, bold in split_bold(text):
        run = paragraph.add_run()
        run.text = segment
        run.font.bold = bold


def _add_bullet_slide(prs: Presentation, title: str, bullets: list[tuple[str, int]]):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, (text, level) in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        _set_bullet_text(p, text, level)


def _add_table_slide(prs: Presentation, title: str, rows: list[list[str]]):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    n_rows, n_cols = len(rows), len(rows[0])
    table = slide.shapes.add_table(
        n_rows, n_cols, Inches(0.5), Inches(1.6), Inches(9), Inches(min(0.5 * n_rows, 5))
    ).table
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            p = table.cell(r, c).text_frame.paragraphs[0]
            for segment, bold in split_bold(val):
                run = p.add_run()
                run.text = segment
                run.font.size = Pt(12)
                run.font.bold = bold or r == 0


def _flush_section(prs: Presentation, section_title: str, bullets: list[tuple[str, int]]):
    for i in range(0, len(bullets), MAX_BULLETS_PER_SLIDE):
        chunk = bullets[i:i + MAX_BULLETS_PER_SLIDE]
        suffix = " (suite)" if i > 0 else ""
        _add_bullet_slide(prs, section_title + suffix, chunk)


def markdown_to_pptx(title: str, sections: list[dict]) -> io.BytesIO:
    prs = Presentation()

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = title
    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = "Account Plan"

    for section in sections:
        bullets: list[tuple[str, int]] = []
        for block in parse_blocks(section["content"]):
            t = block["type"]
            if t == "table":
                if bullets:
                    _flush_section(prs, section["title"], bullets)
                    bullets = []
                _add_table_slide(prs, section["title"], block["rows"])
            elif t in ("h1", "h2", "h3"):
                bullets.append((block["text"], 0))
            elif t in ("bullet", "number"):
                bullets.append((block["text"], 1))
            else:
                bullets.append((block["text"], 0))
        if bullets:
            _flush_section(prs, section["title"], bullets)

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer
