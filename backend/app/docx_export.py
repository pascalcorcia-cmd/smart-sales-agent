import io

from docx import Document

from app.markdown_blocks import parse_blocks, split_bold


def _add_inline_runs(paragraph, text: str):
    for segment, bold in split_bold(text):
        paragraph.add_run(segment).bold = bold or None


def _add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Light Grid Accent 1"
    for r, row in enumerate(rows):
        cells = table.rows[r].cells
        for c, val in enumerate(row):
            if c < len(cells):
                cell_p = cells[c].paragraphs[0]
                _add_inline_runs(cell_p, val)
                if r == 0:
                    for run in cell_p.runs:
                        run.bold = True
    doc.add_paragraph()


def markdown_to_docx(title: str, content: str) -> io.BytesIO:
    doc = Document()
    doc.add_heading(title, level=1)

    for block in parse_blocks(content):
        t = block["type"]
        if t == "table":
            _add_table(doc, block["rows"])
        elif t == "h1":
            doc.add_heading(block["text"], level=1)
        elif t == "h2":
            doc.add_heading(block["text"], level=2)
        elif t == "h3":
            doc.add_heading(block["text"], level=3)
        elif t == "bullet":
            _add_inline_runs(doc.add_paragraph(style="List Bullet"), block["text"])
        elif t == "number":
            _add_inline_runs(doc.add_paragraph(style="List Number"), block["text"])
        else:
            _add_inline_runs(doc.add_paragraph(), block["text"])

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
